import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def create_element(name):
    return OxmlElement(name)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def generate_docx():
    current_dir = os.path.dirname(os.path.abspath(__file__))
    md_path = os.path.join(current_dir, '..', 'Otchet_Kursovaya_Avtoshkola.md')
    out_docx_path = os.path.join(current_dir, '..', '..', 'Otchet_Kursovaya_Avtoshkola.docx')

    if not os.path.exists(md_path):
        print(f"Ошибка: файл {md_path} не найден!")
        return

    with open(md_path, 'r', encoding='utf-8') as f:
        md_text = f.read()

    doc = docx.Document()

    # Настройка полей по ГОСТ 7.32 (Левое: 30 мм, Правое: 15 мм, Верхнее: 20 мм, Нижнее: 20 мм)
    for section in doc.sections:
        section.top_margin = Mm(20)
        section.bottom_margin = Mm(20)
        section.left_margin = Mm(30)
        section.right_margin = Mm(15)

    # Настройка базового стиля «Normal» (Times New Roman, 14 pt, 1.5 интервал, 1.25 см красная строка)
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(14)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)
    normal_style.paragraph_format.line_spacing = 1.5
    normal_style.paragraph_format.space_after = Pt(0)
    normal_style.paragraph_format.space_before = Pt(0)
    normal_style.paragraph_format.first_line_indent = Mm(12.5)
    normal_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    lines = md_text.split('\n')
    in_code_block = False
    code_lines = []

    for line in lines:
        stripped = line.strip()

        # Блоки исходного кода
        if stripped.startswith('```'):
            if in_code_block:
                in_code_block = False
                code_text = '\n'.join(code_lines)
                p = doc.add_paragraph()
                p.paragraph_format.first_line_indent = Mm(0)
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(4)
                run = p.add_run(code_text)
                run.font.name = 'Courier New'
                run.font.size = Pt(10)
                code_lines = []
            else:
                in_code_block = True
                code_lines = []
            continue

        if in_code_block:
            code_lines.append(line)
            continue

        # Заголовки Markdown
        if stripped.startswith('# '):
            title_text = stripped[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Mm(0)
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run(title_text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(16)
            run.bold = True
        elif stripped.startswith('## '):
            title_text = stripped[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Mm(0)
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(title_text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(14)
            run.bold = True
        elif stripped.startswith('### '):
            title_text = stripped[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Mm(12.5)
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(title_text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(14)
            run.bold = True
            run.italic = True
        elif stripped.startswith('* ') or stripped.startswith('- '):
            item_text = stripped[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Mm(12.5)
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.add_run('• ' + item_text)
        elif stripped == '---':
            continue
        elif not stripped:
            continue
        else:
            # Обычный текст
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Mm(12.5)
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.add_run(stripped)

    doc.save(out_docx_path)
    print(f"✅ Отчет по ГОСТ 7.32 успешно скомпилирован в Word: {out_docx_path}")

if __name__ == '__main__':
    generate_docx()
